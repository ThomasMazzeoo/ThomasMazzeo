import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.chart import ScatterChart, Reference, Series

def create_scheda_laboratorio(output_path):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Scheda Esperimento"

    ws.page_setup.orientation = ws.ORIENTATION_LANDSCAPE
    ws.page_setup.paperSize = ws.PAPERSIZE_A4
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 1
    ws.sheet_properties.pageSetUpPr.fitToPage = True

    ws.page_margins.left = 0.35
    ws.page_margins.right = 0.35
    ws.page_margins.top = 0.35
    ws.page_margins.bottom = 0.35
    ws.page_margins.header = 0.15
    ws.page_margins.footer = 0.15

    COLOR_NAVY = "1E3A8A"
    COLOR_BLUE_LIGHT = "DBEAFE"
    COLOR_BLUE_ROW = "EFF6FF"
    COLOR_GREEN_ROW = "ECFDF5"
    COLOR_GREEN_ACCENT = "DCFCE7"
    COLOR_RED_ROW = "FFF1F2"
    COLOR_RED_ACCENT = "FEE2E2"
    COLOR_GRAY_CARD = "F8FAFC"
    COLOR_GRAY_TEXT = "475569"

    font_title = Font(name="Segoe UI", size=13, bold=True, color="0F172A")
    font_subtitle = Font(name="Segoe UI", size=8.5, italic=True, color="475569")
    font_sec_hdr = Font(name="Segoe UI", size=9.5, bold=True, color="1E3A8A")
    font_tbl_hdr = Font(name="Segoe UI", size=8.5, bold=True, color="FFFFFF")
    font_tbl_bold = Font(name="Segoe UI", size=8.5, bold=True, color="1E293B")
    font_footer = Font(name="Segoe UI", size=8, color="1E293B")

    fill_tbl_hdr = PatternFill("solid", fgColor=COLOR_NAVY)
    fill_andata_bg = PatternFill("solid", fgColor=COLOR_BLUE_ROW)
    fill_andata_acc = PatternFill("solid", fgColor=COLOR_BLUE_LIGHT)
    fill_inv_bg = PatternFill("solid", fgColor=COLOR_GREEN_ROW)
    fill_inv_acc = PatternFill("solid", fgColor=COLOR_GREEN_ACCENT)
    fill_rit_bg = PatternFill("solid", fgColor=COLOR_RED_ROW)
    fill_rit_acc = PatternFill("solid", fgColor=COLOR_RED_ACCENT)
    fill_white = PatternFill("solid", fgColor="FFFFFF")
    fill_gray_card = PatternFill("solid", fgColor=COLOR_GRAY_CARD)

    thin_border = Side(style="thin", color="CBD5E1")
    med_navy = Side(style="medium", color="1E3A8A")
    border_cell = Border(left=thin_border, right=thin_border, top=thin_border, bottom=thin_border)
    border_header = Border(left=thin_border, right=thin_border, top=med_navy, bottom=med_navy)

    # Header
    ws.merge_cells("B1:E1")
    ws["B1"] = "ISS ARCHIMEDE — LABORATORIO DI FISICA SPERIMENTALE — 2ª ELETTRONICA"
    ws["B1"].font = Font(name="Segoe UI", size=8.5, bold=True, color="3B82F6")

    ws.merge_cells("B2:E2")
    ws["B2"] = "LABORATORIO 1: IL GRAFICO SPAZIO-TEMPO NEL MOTO RETTILINEO"
    ws["B2"].font = font_title

    ws.merge_cells("B3:E3")
    ws["B3"] = "Esperienza con rotaia e carrello: studio del moto rettilineo con andata e ritorno (punta 2,30 m)"
    ws["B3"].font = font_subtitle

    # Student Info Box (G1:L3)
    ws["G1"] = "Studente/i:"
    ws["G1"].font = Font(name="Segoe UI", size=8.5, bold=True, color="334155")
    ws.merge_cells("H1:J1")
    ws["H1"] = "_______________________________"

    ws["K1"] = "Classe:"
    ws["K1"].font = Font(name="Segoe UI", size=8.5, bold=True, color="334155")
    ws["L1"] = "2ª Elettronica"
    ws["L1"].font = Font(name="Segoe UI", size=8.5, bold=True, color="1E3A8A")

    ws["G2"] = "Docente:"
    ws["G2"].font = Font(name="Segoe UI", size=8.5, bold=True, color="334155")
    ws.merge_cells("H2:J2")
    ws["H2"] = "Prof. Thomas Mazzeo"
    ws["H2"].font = Font(name="Segoe UI", size=8.5, bold=True, color="1E3A8A")

    ws["K2"] = "Data:"
    ws["K2"].font = Font(name="Segoe UI", size=8.5, bold=True, color="334155")
    ws["L2"] = "___ / ___ / 202..."
    ws["L2"].font = Font(name="Segoe UI", size=8.5, color="1E293B")

    # Section Titles
    ws["B5"] = "1. TABELLA DEI DATI SPERIMENTALI (DA COMPILARE)"
    ws["B5"].font = font_sec_hdr

    ws.merge_cells("G5:L5")
    ws["G5"] = "2. GRAFICO SPAZIO-TEMPO s(t) — AUTOMATICO IN TEMPO REALE"
    ws["G5"].font = font_sec_hdr

    # Table Headers
    headers = ["Fase", "Tempo t (s)", "Spazio s (m)", "Note / Evento"]
    cols = ["B", "C", "D", "E"]
    for col, h in zip(cols, headers):
        cell = ws[f"{col}6"]
        cell.value = h
        cell.font = font_tbl_hdr
        cell.fill = fill_tbl_hdr
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = border_header

    # Data Rows: TEMPO t È VUOTO (tranne partenza 0.00) PERCHÉ DEVE ESSERE INSERITO DALLO STUDENTE!
    student_rows = [
        ("Andata", 0.00, 0.00, "Partenza (t0)", fill_andata_bg, fill_andata_acc),
        ("Andata", None, 0.50, "Traguardo 1", fill_andata_bg, fill_andata_acc),
        ("Andata", None, 1.00, "Traguardo 2", fill_andata_bg, fill_andata_acc),
        ("Andata", None, 1.50, "Traguardo 3 (andata)", fill_andata_bg, fill_andata_acc),
        ("Punta",   None, 2.30, "Inversione (punta)", fill_inv_bg, fill_inv_acc),
        ("Ritorno", None, 1.50, "Traguardo 3 (ritorno)", fill_rit_bg, fill_rit_acc),
        ("Ritorno", None, 1.00, "Traguardo 2 (ritorno)", fill_rit_bg, fill_rit_acc),
        ("Ritorno", None, 0.50, "Traguardo 1 (ritorno)", fill_rit_bg, fill_rit_acc),
        ("Ritorno", None, 0.00, "Rientro al punto iniziale", fill_rit_bg, fill_rit_acc),
    ]

    for idx, (fase, t_val, s_val, nota, bg_fill, s_fill) in enumerate(student_rows, start=7):
        c_fase = ws[f"B{idx}"]
        c_fase.value = fase
        c_fase.font = font_tbl_bold
        c_fase.fill = bg_fill
        c_fase.alignment = Alignment(horizontal="center", vertical="center")
        c_fase.border = border_cell

        c_t = ws[f"C{idx}"]
        c_t.value = t_val
        c_t.font = Font(name="Segoe UI", size=9, bold=True, color="000000")
        c_t.fill = fill_white
        c_t.alignment = Alignment(horizontal="center", vertical="center")
        c_t.border = border_cell
        c_t.number_format = "0.00"

        c_s = ws[f"D{idx}"]
        c_s.value = s_val
        c_s.font = font_tbl_bold
        c_s.fill = s_fill
        c_s.alignment = Alignment(horizontal="center", vertical="center")
        c_s.border = border_cell
        c_s.number_format = "0.00"

        c_n = ws[f"E{idx}"]
        c_n.value = nota
        c_n.font = Font(name="Segoe UI", size=7.5, italic=True, color=COLOR_GRAY_TEXT)
        c_n.fill = bg_fill
        c_n.alignment = Alignment(horizontal="left", vertical="center", indent=1)
        c_n.border = border_cell

    # Instructions box under table
    ws.merge_cells("B16:E16")
    ws["B16"] = "✨ GRAFICO AUTOMATICO: Inserisci i tuoi tempi nella colonna C, il grafico si traccia da solo!"
    ws["B16"].font = Font(name="Segoe UI", size=7.5, bold=True, color="1E3A8A")
    ws["B16"].fill = PatternFill("solid", fgColor="EFF6FF")
    ws["B16"].alignment = Alignment(horizontal="left", vertical="center", indent=1)

    ws.merge_cells("B17:E17")
    ws["B17"] = "Punta di inversione a 2,30 m • Il carrello passa due volte per s = 1,50 m (andata e ritorno)."
    ws["B17"].font = Font(name="Segoe UI", size=7.2, color="334155")
    ws["B17"].fill = fill_gray_card
    ws["B17"].alignment = Alignment(horizontal="left", vertical="center", indent=1)

    ws.merge_cells("B18:E18")
    ws["B18"] = "Strumenti: Rotaia graduata con carrello virtuale e cronometro digitale."
    ws["B18"].font = Font(name="Segoe UI", size=7.2, italic=True, color="475569")
    ws["B18"].fill = fill_gray_card
    ws["B18"].alignment = Alignment(horizontal="left", vertical="center", indent=1)

    for r in range(16, 19):
        for c in ["B", "C", "D", "E"]:
            ws[f"{c}{r}"].border = border_cell

    # SCATTER CHART CON LINEE RETTE (NO SPLINE CURVE!)
    chart = ScatterChart()
    chart.title = "Grafico Spazio-Tempo s(t)"
    chart.style = 13
    chart.x_axis.title = "Tempo t [s]"
    chart.y_axis.title = "Spazio s [m]"

    xvalues = Reference(ws, min_col=3, min_row=7, max_row=15)
    yvalues = Reference(ws, min_col=4, min_row=6, max_row=15)

    series = Series(yvalues, xvalues, title_from_data=True)
    series.marker.symbol = "circle"
    series.marker.size = 7
    series.smooth = False  # RIGOROSAMENTE RETTO, NESSUNA CURVA FITTIZIA!
    series.graphicalProperties.line.width = 25400
    series.graphicalProperties.line.solidFill = "1E3A8A"
    series.marker.graphicalProperties.solidFill = "38BDF8"
    series.marker.graphicalProperties.line.solidFill = "1E3A8A"

    chart.series.append(series)
    chart.legend = None
    chart.width = 18.0
    chart.height = 10.4

    ws.add_chart(chart, "G6")

    # Calculations row (19-21) CON FORMULE PROTETTE DA #DIV/0!
    ws.merge_cells("B19:D19")
    ws["B19"] = "Pendenza Andata: v_andata = (2,30 m - 0,00 m)/(t_punta - t_0)"
    ws["B19"].font = font_footer
    ws["E19"] = '=IF(AND(ISNUMBER(C11), ISNUMBER(C7), C11>C7), ROUND((D11-D7)/(C11-C7), 2), "")'
    ws["E19"].font = Font(name="Segoe UI", size=9, bold=True, color="1E3A8A")
    ws["E19"].alignment = Alignment(horizontal="center", vertical="center")
    ws["E19"].number_format = "0.00"

    ws.merge_cells("B20:D20")
    ws["B20"] = "Pendenza Ritorno: v_ritorno = (0,00 m - 2,30 m)/(t_fine - t_punta)"
    ws["B20"].font = font_footer
    ws["E20"] = '=IF(AND(ISNUMBER(C15), ISNUMBER(C11), C15>C11), ROUND((D15-D11)/(C15-C11), 2), "")'
    ws["E20"].font = Font(name="Segoe UI", size=9, bold=True, color="B91C1C")
    ws["E20"].alignment = Alignment(horizontal="center", vertical="center")
    ws["E20"].number_format = "0.00"

    ws.merge_cells("B21:E21")
    ws["B21"] = "Nota: La pendenza di andata è positiva (+), quella di ritorno è negativa (-) perché il verso è opposto."
    ws["B21"].font = Font(name="Segoe UI", size=7.5, italic=True, color="475569")

    # Grade Box
    ws.merge_cells("J19:L19")
    ws["J19"] = "VALUTAZIONE DOCENTE"
    ws["J19"].font = Font(name="Segoe UI", size=8, bold=True, color="1E3A8A")
    ws["J19"].alignment = Alignment(horizontal="center", vertical="center")
    ws["J19"].fill = PatternFill("solid", fgColor="EFF6FF")

    ws.merge_cells("J20:K20")
    ws["J20"] = "Voto: ______ / 10"
    ws["J20"].font = Font(name="Segoe UI", size=8.5, bold=True, color="0F172A")
    ws["J20"].alignment = Alignment(horizontal="center", vertical="center")

    ws["L20"] = "Firma: ________"
    ws["L20"].font = Font(name="Segoe UI", size=7.5, color="475569")
    ws["L20"].alignment = Alignment(horizontal="left", vertical="center")

    for r in range(19, 21):
        for col_l in ["J", "K", "L"]:
            ws[f"{col_l}{r}"].border = border_cell

    # Column Widths
    ws.column_dimensions["A"].width = 2
    ws.column_dimensions["B"].width = 12
    ws.column_dimensions["C"].width = 13
    ws.column_dimensions["D"].width = 13
    ws.column_dimensions["E"].width = 23
    ws.column_dimensions["F"].width = 3
    for col_l in ["G", "H", "I", "J", "K", "L"]:
        ws.column_dimensions[col_l].width = 12

    # Row Heights
    ws.row_dimensions[1].height = 16
    ws.row_dimensions[2].height = 20
    ws.row_dimensions[3].height = 15
    ws.row_dimensions[4].height = 5
    ws.row_dimensions[5].height = 18
    ws.row_dimensions[6].height = 20
    for r in range(7, 16):
        ws.row_dimensions[r].height = 17
    ws.row_dimensions[16].height = 16
    ws.row_dimensions[17].height = 16
    ws.row_dimensions[18].height = 16
    ws.row_dimensions[19].height = 17
    ws.row_dimensions[20].height = 17
    ws.row_dimensions[21].height = 17

    ws.print_area = "A1:L21"
    wb.save(output_path)
    print(f"Scheda Laboratorio salvata con successo in: {output_path}")


def create_scheda_velocita(output_path):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Scheda Velocità"

    ws.page_setup.orientation = ws.ORIENTATION_LANDSCAPE
    ws.page_setup.paperSize = ws.PAPERSIZE_A4
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 1
    ws.sheet_properties.pageSetUpPr.fitToPage = True

    ws.page_margins.left = 0.35
    ws.page_margins.right = 0.35
    ws.page_margins.top = 0.35
    ws.page_margins.bottom = 0.35
    ws.page_margins.header = 0.15
    ws.page_margins.footer = 0.15

    COLOR_NAVY = "1E3A8A"
    COLOR_BLUE_LIGHT = "DBEAFE"
    COLOR_BLUE_ROW = "EFF6FF"
    COLOR_GREEN_ROW = "ECFDF5"
    COLOR_GREEN_ACCENT = "DCFCE7"
    COLOR_RED_ROW = "FFF1F2"
    COLOR_RED_ACCENT = "FEE2E2"
    COLOR_GRAY_CARD = "F8FAFC"
    COLOR_GRAY_TEXT = "475569"

    font_title = Font(name="Segoe UI", size=13, bold=True, color="0F172A")
    font_subtitle = Font(name="Segoe UI", size=8.5, italic=True, color="475569")
    font_sec_hdr = Font(name="Segoe UI", size=9.5, bold=True, color="1E3A8A")
    font_tbl_hdr = Font(name="Segoe UI", size=8.5, bold=True, color="FFFFFF")
    font_tbl_bold = Font(name="Segoe UI", size=8.5, bold=True, color="1E293B")
    font_footer = Font(name="Segoe UI", size=8, color="1E293B")

    fill_tbl_hdr = PatternFill("solid", fgColor=COLOR_NAVY)
    fill_andata_bg = PatternFill("solid", fgColor=COLOR_BLUE_ROW)
    fill_inv_bg = PatternFill("solid", fgColor=COLOR_GREEN_ROW)
    fill_rit_bg = PatternFill("solid", fgColor=COLOR_RED_ROW)
    fill_white = PatternFill("solid", fgColor="FFFFFF")
    fill_gray_card = PatternFill("solid", fgColor=COLOR_GRAY_CARD)

    thin_border = Side(style="thin", color="CBD5E1")
    med_navy = Side(style="medium", color="1E3A8A")
    border_cell = Border(left=thin_border, right=thin_border, top=thin_border, bottom=thin_border)
    border_header = Border(left=thin_border, right=thin_border, top=med_navy, bottom=med_navy)

    # Header
    ws.merge_cells("B1:E1")
    ws["B1"] = "ISS ARCHIMEDE — LABORATORIO DI FISICA SPERIMENTALE — 2ª ELETTRONICA"
    ws["B1"].font = Font(name="Segoe UI", size=8.5, bold=True, color="3B82F6")

    ws.merge_cells("B2:E2")
    ws["B2"] = "IL GRAFICO VELOCITÀ-TEMPO NEL MOTO RETTILINEO"
    ws["B2"].font = font_title

    ws.merge_cells("B3:E3")
    ws["B3"] = "Studio della velocità media e rappresentazione del grafico v(t) con andata e ritorno"
    ws["B3"].font = font_subtitle

    # Student Info Box (G1:L3)
    ws["G1"] = "Studente/i:"
    ws["G1"].font = Font(name="Segoe UI", size=8.5, bold=True, color="334155")
    ws.merge_cells("H1:J1")
    ws["H1"] = "_______________________________"

    ws["K1"] = "Classe:"
    ws["K1"].font = Font(name="Segoe UI", size=8.5, bold=True, color="334155")
    ws["L1"] = "2ª Elettronica"
    ws["L1"].font = Font(name="Segoe UI", size=8.5, bold=True, color="1E3A8A")

    ws["G2"] = "Docente:"
    ws["G2"].font = Font(name="Segoe UI", size=8.5, bold=True, color="334155")
    ws.merge_cells("H2:J2")
    ws["H2"] = "Prof. Thomas Mazzeo"
    ws["H2"].font = Font(name="Segoe UI", size=8.5, bold=True, color="1E3A8A")

    ws["K2"] = "Data:"
    ws["K2"].font = Font(name="Segoe UI", size=8.5, bold=True, color="334155")
    ws["L2"] = "___ / ___ / 202..."
    ws["L2"].font = Font(name="Segoe UI", size=8.5, color="1E293B")

    # Section Titles
    ws["B5"] = "1. TABELLA DEI DATI CALCOLATI (AUTOMATICA)"
    ws["B5"].font = font_sec_hdr

    ws.merge_cells("G5:L5")
    ws["G5"] = "2. GRAFICO VELOCITÀ-TEMPO v(t) — AUTOMATICO IN TEMPO REALE"
    ws["G5"].font = font_sec_hdr

    # Table Headers
    headers = ["Fase", "Tempo t (s)", "Velocità v (m/s)", "Note / Traguardo"]
    cols = ["B", "C", "D", "E"]
    for col, h in zip(cols, headers):
        cell = ws[f"{col}6"]
        cell.value = h
        cell.font = font_tbl_hdr
        cell.fill = fill_tbl_hdr
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = border_header

    # Data Rows: I TEMPI C SONO VUOTI (lo studente li inserisce). LA VELOCITÀ D È UNA FORMULA!
    vel_specs = [
        # (fase, t_val, formula_vel, nota, bg_fill)
        ("Andata", 0.00, 0.00, "Partenza (t0)", fill_andata_bg),
        ("Andata", None, '=IF(AND(ISNUMBER(C8), ISNUMBER(C7), C8>C7), ROUND(0.50/(C8-C7), 2), "")', "Traguardo 1 (0,50 m)", fill_andata_bg),
        ("Andata", None, '=IF(AND(ISNUMBER(C9), ISNUMBER(C8), C9>C8), ROUND(0.50/(C9-C8), 2), "")', "Traguardo 2 (1,00 m)", fill_andata_bg),
        ("Andata", None, '=IF(AND(ISNUMBER(C10), ISNUMBER(C9), C10>C9), ROUND(0.50/(C10-C9), 2), "")', "Traguardo 3 (1,50 m)", fill_andata_bg),
        ("Punta",   None, '=IF(AND(ISNUMBER(C11), ISNUMBER(C10), C11>C10), ROUND(0.80/(C11-C10), 2), "")', "Inversione (2,30 m)", fill_inv_bg),
        ("Ritorno", None, '=IF(AND(ISNUMBER(C12), ISNUMBER(C11), C12>C11), ROUND(-0.80/(C12-C11), 2), "")', "Traguardo 3 rit. (1,50 m)", fill_rit_bg),
        ("Ritorno", None, '=IF(AND(ISNUMBER(C13), ISNUMBER(C12), C13>C12), ROUND(-0.50/(C13-C12), 2), "")', "Traguardo 2 rit. (1,00 m)", fill_rit_bg),
        ("Ritorno", None, '=IF(AND(ISNUMBER(C14), ISNUMBER(C13), C14>C13), ROUND(-0.50/(C14-C13), 2), "")', "Traguardo 1 rit. (0,50 m)", fill_rit_bg),
        ("Ritorno", None, '=IF(AND(ISNUMBER(C15), ISNUMBER(C14), C15>C14), ROUND(-0.50/(C15-C14), 2), "")', "Rientro alla partenza (0,00 m)", fill_rit_bg),
    ]

    for idx, (fase, t_val, v_formula, nota, bg_fill) in enumerate(vel_specs, start=7):
        c_fase = ws[f"B{idx}"]
        c_fase.value = fase
        c_fase.font = font_tbl_bold
        c_fase.fill = bg_fill
        c_fase.alignment = Alignment(horizontal="center", vertical="center")
        c_fase.border = border_cell

        c_t = ws[f"C{idx}"]
        c_t.value = t_val
        c_t.font = Font(name="Segoe UI", size=9, bold=True, color="000000")
        c_t.fill = fill_white
        c_t.alignment = Alignment(horizontal="center", vertical="center")
        c_t.border = border_cell
        c_t.number_format = "0.00"

        c_v = ws[f"D{idx}"]
        c_v.value = v_formula
        c_v.font = Font(name="Segoe UI", size=9, bold=True, color="1E3A8A")
        c_v.fill = fill_white
        c_v.alignment = Alignment(horizontal="center", vertical="center")
        c_v.border = border_cell
        c_v.number_format = "+0.00;-0.00;0.00"

        c_n = ws[f"E{idx}"]
        c_n.value = nota
        c_n.font = Font(name="Segoe UI", size=7.5, italic=True, color=COLOR_GRAY_TEXT)
        c_n.fill = bg_fill
        c_n.alignment = Alignment(horizontal="left", vertical="center", indent=1)
        c_n.border = border_cell

    # Instructions box under table
    ws.merge_cells("B16:E16")
    ws["B16"] = "✨ GRAFICO AUTOMATICO: Inserisci i tempi in colonna C, velocità e grafico v(t) si calcolano da soli!"
    ws["B16"].font = Font(name="Segoe UI", size=7.5, bold=True, color="1E3A8A")
    ws["B16"].fill = PatternFill("solid", fgColor="EFF6FF")
    ws["B16"].alignment = Alignment(horizontal="left", vertical="center", indent=1)

    ws.merge_cells("B17:E17")
    ws["B17"] = "Andata: velocità positiva v > 0 (+) • Ritorno: velocità negativa v < 0 (-) verso opposto."
    ws["B17"].font = Font(name="Segoe UI", size=7.2, color="334155")
    ws["B17"].fill = fill_gray_card
    ws["B17"].alignment = Alignment(horizontal="left", vertical="center", indent=1)

    ws.merge_cells("B18:E18")
    ws["B18"] = "Strumenti: Rotaia a cuscino d'aria con carrello e cronometro digitale."
    ws["B18"].font = Font(name="Segoe UI", size=7.2, italic=True, color="475569")
    ws["B18"].fill = fill_gray_card
    ws["B18"].alignment = Alignment(horizontal="left", vertical="center", indent=1)

    for r in range(16, 19):
        for c in ["B", "C", "D", "E"]:
            ws[f"{c}{r}"].border = border_cell

    # NATIVE EXCEL SCATTER CHART (Velocity-Time!)
    # Nel MRU: niente curve spline! Usiamo linee rette o marcatori discreti senza smoothing!
    chart = ScatterChart()
    chart.title = "Grafico Velocità-Tempo v(t)"
    chart.style = 13
    chart.x_axis.title = "Tempo t [s]"
    chart.y_axis.title = "Velocità v [m/s]"

    xvalues = Reference(ws, min_col=3, min_row=7, max_row=15)
    yvalues = Reference(ws, min_col=4, min_row=6, max_row=15)

    series = Series(yvalues, xvalues, title_from_data=True)
    series.marker.symbol = "square"
    series.marker.size = 7
    series.smooth = False  # NO ONDE CURVE O SPLINE ARTIFICIALI!
    series.graphicalProperties.line.width = 25400
    series.graphicalProperties.line.solidFill = "16A34A"
    series.marker.graphicalProperties.solidFill = "22C55E"
    series.marker.graphicalProperties.line.solidFill = "16A34A"

    chart.series.append(series)
    chart.legend = None
    chart.width = 18.0
    chart.height = 10.4

    ws.add_chart(chart, "G6")

    # Overall Average Velocities
    ws.merge_cells("B19:D19")
    ws["B19"] = "Velocità Media Andata (tratto 0 m -> 2,30 m):"
    ws["B19"].font = font_footer
    ws["E19"] = '=IF(AND(ISNUMBER(C11), C11>0), ROUND(2.30/C11, 2), "")'
    ws["E19"].font = Font(name="Segoe UI", size=9, bold=True, color="16A34A")
    ws["E19"].alignment = Alignment(horizontal="center", vertical="center")
    ws["E19"].number_format = "+0.00;-0.00;0.00"

    ws.merge_cells("B20:D20")
    ws["B20"] = "Velocità Media Ritorno (tratto 2,30 m -> 0 m):"
    ws["B20"].font = font_footer
    ws["E20"] = '=IF(AND(ISNUMBER(C15), ISNUMBER(C11), C15>C11), ROUND(-2.30/(C15-C11), 2), "")'
    ws["E20"].font = Font(name="Segoe UI", size=9, bold=True, color="DC2626")
    ws["E20"].alignment = Alignment(horizontal="center", vertical="center")
    ws["E20"].number_format = "+0.00;-0.00;0.00"

    ws.merge_cells("B21:E21")
    ws["B21"] = "Domanda: Perché il grafico della velocità cambia segno a metà percorso? Perché il moto inverte il verso di marcia."
    ws["B21"].font = Font(name="Segoe UI", size=7.5, italic=True, color="475569")

    # Grade Box
    ws.merge_cells("J19:L19")
    ws["J19"] = "VALUTAZIONE DOCENTE"
    ws["J19"].font = Font(name="Segoe UI", size=8, bold=True, color="1E3A8A")
    ws["J19"].alignment = Alignment(horizontal="center", vertical="center")
    ws["J19"].fill = PatternFill("solid", fgColor="EFF6FF")

    ws.merge_cells("J20:K20")
    ws["J20"] = "Voto: ______ / 10"
    ws["J20"].font = Font(name="Segoe UI", size=8.5, bold=True, color="0F172A")
    ws["J20"].alignment = Alignment(horizontal="center", vertical="center")

    ws["L20"] = "Firma: ________"
    ws["L20"].font = Font(name="Segoe UI", size=7.5, color="475569")
    ws["L20"].alignment = Alignment(horizontal="left", vertical="center")

    for r in range(19, 21):
        for col_l in ["J", "K", "L"]:
            ws[f"{col_l}{r}"].border = border_cell

    # Column Widths
    ws.column_dimensions["A"].width = 2
    ws.column_dimensions["B"].width = 12
    ws.column_dimensions["C"].width = 13
    ws.column_dimensions["D"].width = 13
    ws.column_dimensions["E"].width = 23
    ws.column_dimensions["F"].width = 3
    for col_l in ["G", "H", "I", "J", "K", "L"]:
        ws.column_dimensions[col_l].width = 12

    # Row Heights
    ws.row_dimensions[1].height = 16
    ws.row_dimensions[2].height = 20
    ws.row_dimensions[3].height = 15
    ws.row_dimensions[4].height = 5
    ws.row_dimensions[5].height = 18
    ws.row_dimensions[6].height = 20
    for r in range(7, 16):
        ws.row_dimensions[r].height = 17
    ws.row_dimensions[16].height = 16
    ws.row_dimensions[17].height = 16
    ws.row_dimensions[18].height = 16
    ws.row_dimensions[19].height = 17
    ws.row_dimensions[20].height = 17
    ws.row_dimensions[21].height = 17

    ws.print_area = "A1:L21"
    wb.save(output_path)
    print(f"Scheda Velocità salvata con successo in: {output_path}")

if __name__ == "__main__":
    targets = [
        ("public/fisica/archimede/seconda-elettronica/scheda_laboratorio_A4.xlsx", "public/fisica/archimede/seconda-elettronica/scheda_velocita_A4.xlsx"),
        ("C:/Users/thoma/Desktop/LAB FISICA/ESPERIMENTO SECONDA ELETTRONICA 1/scheda_laboratorio_A4.xlsx", "C:/Users/thoma/Desktop/LAB FISICA/ESPERIMENTO SECONDA ELETTRONICA 1/scheda_velocita_A4.xlsx"),
    ]
    for p_lab, p_vel in targets:
        create_scheda_laboratorio(p_lab)
        create_scheda_velocita(p_vel)
    print("Tutti i file Excel generati con successo!")
